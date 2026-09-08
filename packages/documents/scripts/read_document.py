"""Bounded, read-only document extraction. Input content is always untrusted data."""
from pathlib import Path
import csv
import io
import json
import re
import sys
import zipfile
from lxml import etree

LIMIT = 60000
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

def package(path):
    archive = zipfile.ZipFile(path)
    names = archive.namelist()
    if len(names) != len(set(names)):
        archive.close()
        raise ValueError("The Office package contains duplicate parts and requires review.")
    if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
        archive.close()
        raise ValueError("The Office package exceeds the preview limit.")
    return archive

def xml(archive, name):
    return etree.fromstring(archive.read(name), etree.XMLParser(resolve_entities=False, no_network=True))

def preview_row(values):
    """Keep the clipping acknowledgement aligned with every displayed value."""
    truncated = len(values) > 30
    row = []
    for value in values[:30]:
        text = "" if value is None else str(value)
        truncated = truncated or len(text) > 500
        row.append(text[:500])
    return row, truncated

def verify(path):
    if path.with_name("~$" + path.name).exists():
        raise ValueError("A Word owner file is present. Close the document before previewing it.")
    with package(path) as archive:
        for name in archive.namelist():
            if name.endswith(".rels"):
                for rel in xml(archive, name):
                    if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                        raise ValueError("This document links to external resources. Download it for review.")
            if not name.endswith(".xml"):
                continue
            for element in xml(archive, name).iter():
                for attribute, value in element.attrib.items():
                    if attribute in {"{" + MC + "}" + key for key in ("Ignorable", "ProcessContent", "PreserveElements", "PreserveAttributes")} or (element.tag == "{" + MC + "}Choice" and attribute == "Requires"):
                        for item in value.split():
                            prefix = item.split(":")[0]
                            if prefix not in element.nsmap:
                                raise ValueError("The document references an undeclared compatibility namespace.")
        if any("vbaproject" in name.lower() for name in archive.namelist()):
            raise ValueError("Macro-enabled packages are not previewed.")
    return {"verified": True}

def extract(path):
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        with path.open("rb") as handle:
            raw = handle.read(LIMIT * 4 + 1)
        text = raw.decode("utf-8", errors="replace")
        return {"kind": "text", "text": text[:LIMIT], "truncated": len(text) > LIMIT}
    if ext in {".xlsx", ".csv"}:
        sheets = []
        truncated = False
        if ext == ".xlsx":
            import openpyxl
            with package(path):
                pass
            book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book.worksheets[:8]:
                    rows = []
                    for values in sheet.iter_rows(max_row=min(sheet.max_row or 1, 200), max_col=min(sheet.max_column or 1, 30), values_only=True):
                        row, clipped = preview_row(values)
                        rows.append(row)
                        truncated = truncated or clipped
                    sheets.append({"name": sheet.title, "rows": rows})
                    truncated = truncated or (sheet.max_row or 0) > 200 or (sheet.max_column or 0) > 30
                truncated = truncated or len(book.worksheets) > 8
            finally:
                book.close()
        else:
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                rows = []
                for index, row in enumerate(csv.reader(handle)):
                    if index >= 200:
                        truncated = True
                        break
                    visible, clipped = preview_row(row)
                    rows.append(visible)
                    truncated = truncated or clipped
            sheets.append({"name": "Sheet", "rows": rows})
        return {"kind": "spreadsheet", "sheets": sheets, "truncated": truncated, "message": "Read-only cached values. Formulas and external links are not recalculated."}
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("This PDF requires a password.")
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages[:30])
        return {"kind": "text", "text": text[:LIMIT], "truncated": len(text) > LIMIT or len(reader.pages) > 30}
    if ext in {".docx", ".pptx"}:
        with package(path) as archive:
            if ext == ".docx":
                names = [name for name in archive.namelist() if re.fullmatch(r"word/(document|header\d+|footer\d+)\.xml", name)]
                texts = ["\n".join(
                    "".join((node.text or "") if node.tag.endswith("}t") else ("\t" if node.tag.endswith("}tab") else "\n" if node.tag.endswith("}br") else "") for node in paragraph.iter())
                    for paragraph in xml(archive, name).iter() if paragraph.tag.endswith("}p")
                ) for name in names]
            else:
                names = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)), key=lambda name: int(re.search(r"slide(\d+)", name)[1]))
                texts = ["\n".join(element.text or "" for element in xml(archive, name).iter() if element.tag.endswith("}t")) for name in names[:100]]
            text = "\n\n".join(texts)
        return {"kind": "text", "text": text[:LIMIT], "truncated": len(text) > LIMIT}
    return {"kind": "unsupported", "message": "Download the original to open this format."}

if __name__ == "__main__":
    try:
        mode, filename = sys.argv[1:3]
        result = verify(Path(filename)) if mode == "verify" else extract(Path(filename))
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        # Avoid including source paths or document contents in process diagnostics.
        message = str(error) if isinstance(error, ValueError) else "The document could not be read safely. Download the original for review."
        print(json.dumps({"kind": "failed", "message": message}))
        sys.exit(2)
